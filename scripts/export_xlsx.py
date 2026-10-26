"""Собирает Perechen_SI.xlsx из data/ и исходного source/perechen.xlsx, не сбрасывая оформление."""

import json
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
META = ROOT / "data" / "meta.json"
SOURCE = ROOT / "source" / "perechen.xlsx"
OUT = ROOT / "Perechen_SI.xlsx"

FIELDS = {
    "seq": "№ п/п",
    "name": "Наименование",
    "type": "Тип",
    "serial": "Заводской номер",
    "year": "Год выпуска",
    "range": "Диапазон измерений",
    "accuracy": "Класс точности",
    "location": "Место установки СИ",
    "object": "Объект установки",
    "shop": "Подразделение",
    "installDate": "Дата установки",
    "mpiRaw": "Межповерочный интервал, мес.",
    "lastCheck": "Дата последней поверки / калибровки",
    "checkKind": "Поверка (П) / Калибровка (К)",
    "note": "Примечание",
}


def text(value) -> str:
    return "" if value is None else str(value).strip()


def load_instruments(meta: dict) -> dict[str, list]:
    out = {}
    for name, path in (meta.get("sheetFiles") or {}).items():
        file = ROOT / path
        if not file.exists():
            continue
        data = json.loads(file.read_text(encoding="utf-8"))
        out[name] = data.get("instruments") or []
    return out


def value_of(inst: dict, key: str) -> str:
    if key == "seq":
        return "" if not inst.get("seq") else str(inst.get("seq"))
    return text(inst.get(key))


def write_row(ws, row: int, inst: dict, column_map: dict, template_row: int | None):
    for key, col in column_map.items():
        if not isinstance(col, int):
            continue
        cell = ws.cell(row, col + 1, value_of(inst, key))
        if template_row and template_row != row:
            src = ws.cell(template_row, col + 1)
            if src.has_style:
                cell.font = copy(src.font)
                cell.border = copy(src.border)
                cell.fill = copy(src.fill)
                cell.number_format = src.number_format
                cell.alignment = copy(src.alignment)


def ensure_sheet(wb, name: str, headers: list[str]):
    if name in wb.sheetnames:
        return wb[name]
    ws = wb.create_sheet(name)
    for i, header in enumerate(headers, start=1):
        ws.cell(1, i, header)
    return ws


def main() -> None:
    if not META.exists():
        print("data/meta.json ещё нет, Excel не собран")
        return
    meta = json.loads(META.read_text(encoding="utf-8"))
    groups = load_instruments(meta)
    wb = load_workbook(SOURCE) if SOURCE.exists() else Workbook()
    if not SOURCE.exists() and "Sheet" in wb.sheetnames and len(wb.sheetnames) == 1:
        wb.remove(wb["Sheet"])
    sheets = {sheet["name"]: sheet for sheet in meta.get("sheets") or []}
    for name, instruments in groups.items():
        info = sheets.get(name) or {}
        headers = info.get("headers") or list(FIELDS.values())
        column_map = info.get("columnMap") or {}
        if not column_map:
            column_map = {key: headers.index(label) for key, label in FIELDS.items() if label in headers}
        ws = ensure_sheet(wb, name, headers)
        header_row = int(info.get("headerRow") or 0) + 1
        serial_col = column_map.get("serial")
        by_serial: dict[str, int] = {}
        if isinstance(serial_col, int):
            for row in range(header_row + 1, ws.max_row + 1):
                serial = text(ws.cell(row, serial_col + 1).value)
                if serial:
                    by_serial[serial] = row
        next_row = max(ws.max_row, header_row) + 1
        seq_col = column_map.get("seq")
        seq = 0
        for inst in instruments:
            seq += 1
            inst = {**inst, "seq": seq}
            serial = text(inst.get("serial"))
            row = by_serial.get(serial) if serial else None
            if row is None:
                row = next_row
                next_row += 1
            template = row - 1 if row > header_row + 1 else None
            write_row(ws, row, inst, column_map, template)
            if isinstance(seq_col, int):
                ws.cell(row, seq_col + 1, seq)
    wb.save(OUT)
    print(f"собран {OUT.name}")


if __name__ == "__main__":
    main()
